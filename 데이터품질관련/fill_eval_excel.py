# -*- coding: utf-8 -*-
"""
평가 엑셀 자동 작성 v2 — 질문 16개 + 정답(사람이 읽는 문장) + 정답 SQL

사용법:
  conda activate sbiz
  cd C:\\Users\\USER\\bigdata_sbiz\\RAG\\데이터품질관련
  python fill_eval_excel.py

출력: 상권질의_평가_정답포함.xlsx
  A 예시유형 / B 예시질의 / C llm답변(비움) / D 정답(문장) / E 정답SQL / F 결측·판정기준
비교법: D↔데모 답변(수치), E↔데모 화면의 '생성된 SQL'(필터·로직)
"""
import os
import psycopg2
import psycopg2.extras
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

DSN = os.environ.get("SBIZ_DSN", "postgresql://postgres:postgres@localhost:5432/postgres")
OUT = "상권질의_평가_정답포함.xlsx"


# ── 정답 문장 만들기 도우미 ──────────────────────────────────────────
def won(v):    return f"{v:,.0f}만원" if v is not None else "정보없음"
def cnt(v):    return f"{v:,.0f}개" if v is not None else "정보없음"
def pop(v):    return f"{v:,.0f}명" if v is not None else "정보없음"
def pct(v):    return f"{v:+.1f}%" if v is not None else "정보없음"
def rank_list(rows, fmt_one):
    return " / ".join(f"{i}위 {fmt_one(r)}" for i, r in enumerate(rows, 1))


# ── 질문 정의: (유형, 질의, 정답SQL, 문장변환함수, 판정기준) ─────────
QUESTIONS = [
 ("단순조회", "2026년 3월 기준으로 명동에서 치킨집 하는데 요즘 어때?",
  """SELECT amt_dong, amt_sgg, amt_on_year, upso_count_dong, pop_month_avg_cnt
     FROM clean.simple_analysis
     WHERE admi_nm='명동' AND tpbiz_clscd_nm='치킨' AND grain='dong' AND crtr_ym='202603'""",
  lambda R: (lambda r: f"명동 치킨 월평균 매출 {won(r['amt_dong'])} (중구 평균 {won(r['amt_sgg'])}), "
             f"전년동월 대비 {pct(r['amt_on_year'])}, 업소 {cnt(r['upso_count_dong'])}, "
             f"일평균 유동인구 {pop(r['pop_month_avg_cnt'])}")(R[0]),
  "매출·전년비·업소수·유동인구 4개 수치 일치 여부"),

 ("단순조회", "3월 기준 을지로동 치킨 매출이랑 전년 대비 어때?",
  """SELECT amt_dong, amt_on_year, amt_on_month FROM clean.simple_analysis
     WHERE admi_nm='을지로동' AND tpbiz_clscd_nm='치킨' AND grain='dong' AND crtr_ym='202603'""",
  lambda R: (lambda r: f"을지로동 치킨 월평균 매출 {won(r['amt_dong'])}, "
             f"전년동월 대비 {pct(r['amt_on_year'])}, 전월 대비 {pct(r['amt_on_month'])}")(R[0]),
  ""),

 ("랭킹", "3월 기준 서울에서 치킨 장사 제일 잘되는 동네 3곳은?",
  """SELECT sgg_nm, admi_nm, amt_dong FROM clean.simple_analysis
     WHERE sido_nm_std='서울특별시' AND tpbiz_clscd_nm='치킨' AND grain='dong'
       AND is_substituted=false AND crtr_ym='202603'
     ORDER BY amt_dong DESC NULLS LAST LIMIT 3""",
  lambda R: rank_list(R, lambda r: f"{r['sgg_nm']} {r['admi_nm']} {won(r['amt_dong'])}"),
  "SQL에 grain='dong'과 is_substituted=false 필터 포함 여부도 채점"),

 ("랭킹(전국)", "전국에서 카페 매출 1등 동네는 어디야?",
  """SELECT crtr_ym, sido_nm_std, sgg_nm, admi_nm, amt_dong FROM clean.simple_analysis
     WHERE tpbiz_clscd_nm='카페' AND grain='dong' AND is_substituted=false
       AND crtr_ym=(SELECT max(crtr_ym) FROM clean.simple_analysis)
     ORDER BY amt_dong DESC NULLS LAST LIMIT 1""",
  lambda R: (lambda r: f"{r['sido_nm_std']} {r['sgg_nm']} {r['admi_nm']} {won(r['amt_dong'])} "
             f"({r['crtr_ym'][:4]}년 {int(r['crtr_ym'][4:])}월 기준)")(R[0]),
  "월 미지정 → 최신월(202605) 기준으로 답해야 정상"),

 ("업종비교", "3월 기준 구의1동에서 치킨이랑 카페 중 뭐가 나아?",
  """SELECT tpbiz_clscd_nm, amt_dong, upso_count_dong FROM clean.simple_analysis
     WHERE admi_nm='구의1동' AND tpbiz_clscd_nm IN ('치킨','카페') AND grain='dong' AND crtr_ym='202603'
     ORDER BY tpbiz_clscd_nm""",
  lambda R: " vs ".join(f"{r['tpbiz_clscd_nm']} 매출 {won(r['amt_dong'])}·업소 {cnt(r['upso_count_dong'])}" for r in R),
  "두 업종 매출+업소수를 함께 비교했는지"),

 ("역질의", "3월 기준 서울에서 유동인구 10만 넘는데 치킨집 3개 이하인 동네는?",
  """SELECT sgg_nm, admi_nm, pop_month_avg_cnt, upso_count_dong FROM clean.simple_analysis
     WHERE sido_nm_std='서울특별시' AND tpbiz_clscd_nm='치킨' AND grain='dong'
       AND pop_month_avg_cnt>=100000 AND upso_count_dong BETWEEN 1 AND 3 AND crtr_ym='202603'
     ORDER BY pop_month_avg_cnt DESC LIMIT 5""",
  lambda R: " / ".join(f"{r['sgg_nm']} {r['admi_nm']} (유동 {pop(r['pop_month_avg_cnt'])}, 치킨 {cnt(r['upso_count_dong'])})" for r in R),
  "업소 1~3 조건 질문이라 is_substituted 필터는 없어야 정상"),

 ("하락랭킹", "3월 기준 서울에서 작년보다 치킨 장사가 제일 나빠진 동네 3곳은?",
  """SELECT sgg_nm, admi_nm, amt_on_year FROM clean.simple_analysis
     WHERE sido_nm_std='서울특별시' AND tpbiz_clscd_nm='치킨' AND grain='dong'
       AND is_substituted=false AND amt_on_year IS NOT NULL AND crtr_ym='202603'
     ORDER BY amt_on_year ASC LIMIT 3""",
  lambda R: rank_list(R, lambda r: f"{r['sgg_nm']} {r['admi_nm']} {pct(r['amt_on_year'])}"),
  "하락 = 오름차순 정렬, NULL 제외 여부"),

 ("파생계산", "3월 기준 서울에서 카페 가게 하나당 벌이가 제일 좋은 동네는? (업소 5개 이상)",
  """SELECT sgg_nm, admi_nm, ROUND(amt_dong/NULLIF(upso_count_dong,0),1) AS per_store
     FROM clean.simple_analysis
     WHERE sido_nm_std='서울특별시' AND tpbiz_clscd_nm='카페' AND grain='dong'
       AND is_substituted=false AND upso_count_dong>=5 AND crtr_ym='202603'
     ORDER BY per_store DESC NULLS LAST LIMIT 3""",
  lambda R: rank_list(R, lambda r: f"{r['sgg_nm']} {r['admi_nm']} 업소당 {won(r['per_store'])}"),
  "매출÷업소수 계산을 SQL에서 했는지"),

 ("배달", "3월 기준 20대 상대로 카페 배달하려면 서울 어디가 좋아?",
  """SELECT sgg_nm, admi_nm, baemin_age_20 FROM clean.simple_analysis
     WHERE sido_nm_std='서울특별시' AND tpbiz_clscd_nm='카페' AND grain='dong'
       AND baemin_age_20 IS NOT NULL AND crtr_ym='202603'
     ORDER BY baemin_age_20 DESC LIMIT 3""",
  lambda R: rank_list(R, lambda r: f"{r['sgg_nm']} {r['admi_nm']} (20대 비중 {r['baemin_age_20']:.1f}%)"),
  "배달 컬럼 NULL 제외 여부"),

 ("월추이", "명동 치킨 매출이 3월부터 5월까지 어떻게 변했어?",
  """SELECT crtr_ym, amt_dong FROM clean.simple_analysis
     WHERE admi_nm='명동' AND tpbiz_clscd_nm='치킨' AND grain='dong' ORDER BY crtr_ym""",
  lambda R: " → ".join(f"{int(r['crtr_ym'][4:])}월 {won(r['amt_dong'])}" for r in R),
  "3개월 값 + 증감 방향 서술"),

 ("표기함정", "경기도에서 편의점 매출 제일 높은 동네는?",
  """SELECT sgg_nm, admi_nm, amt_dong FROM clean.simple_analysis
     WHERE sido_nm_std='경기도' AND tpbiz_clscd_nm='편의점' AND grain='dong'
       AND is_substituted=false AND crtr_ym=(SELECT max(crtr_ym) FROM clean.simple_analysis)
     ORDER BY amt_dong DESC NULLS LAST LIMIT 3""",
  lambda R: rank_list(R, lambda r: f"{r['sgg_nm']} {r['admi_nm']} {won(r['amt_dong'])}") + " (최신월 기준)",
  "핵심: 생성 SQL이 sido_nm_std='경기도'를 쓰는지 (sido_nm이면 0건 나옴)"),

 ("동명함정", "중구에서 치킨 장사 잘되는 동네는?",
  """SELECT sido_nm_std, sgg_nm, admi_nm, amt_dong FROM clean.simple_analysis
     WHERE sgg_nm='중구' AND tpbiz_clscd_nm='치킨' AND grain='dong'
       AND is_substituted=false AND crtr_ym=(SELECT max(crtr_ym) FROM clean.simple_analysis)
     ORDER BY amt_dong DESC NULLS LAST LIMIT 5""",
  lambda R: "'중구'는 여러 시도에 존재 — 시도 명시 또는 되묻기가 이상적. 참고(전체 중구 상위): "
            + " / ".join(f"{r['sido_nm_std']} {r['admi_nm']} {won(r['amt_dong'])}" for r in R),
  "특정 시도의 중구로 단정하면 감점, 시도를 밝히거나 되물으면 정답"),

 ("합계이해", "3월 기준 서울 전체 치킨집은 몇 개야?",
  """SELECT DISTINCT upso_count_sido FROM clean.simple_analysis
     WHERE sido_nm_std='서울특별시' AND tpbiz_clscd_nm='치킨' AND crtr_ym='202603'
       AND upso_count_sido IS NOT NULL""",
  lambda R: f"서울 치킨 업소 총 {cnt(R[0]['upso_count_sido'])} (2026년 3월 기준)",
  "upso_count_sido 활용 또는 동별 합산 — 유사값이면 인정"),

 ("거부(개별점포)", "명동에서 제일 잘 버는 치킨집 이름이 뭐야?", None,
  lambda R: "[답변 거부가 정답] 개별 점포 정보 없음(행정동×업종 집계 데이터) — 지역 평균으로 안내해야 함",
  "SQL을 만들어내면 실패 (환각)"),

 ("거부(예측)", "내년에 명동에서 치킨집 차리면 잘될까?", None,
  lambda R: "[답변 거부가 정답] 예측·보장 불가 — 현황 수치(매출·추이)만 안내하면 가점",
  "미래 전망을 단정하면 실패"),

 ("결측처리", "3월 기준 명동 부동산 중개업소 배달 주문 얼마나 돼?",
  """SELECT tpbiz_clscd_nm, baemin_month_avg_cnt FROM clean.simple_analysis
     WHERE admi_nm='명동' AND tpbiz_lclcd_nm='부동산' AND grain='dong' AND crtr_ym='202603'
       AND baemin_month_avg_cnt IS NOT NULL LIMIT 5""",
  lambda R: ("'배달 데이터 없음' 안내가 정답 (부동산 업종은 배달 미발생)" if not R
             else "예외: 값 존재 — " + ", ".join(f"{r['tpbiz_clscd_nm']} {r['baemin_month_avg_cnt']}건" for r in R)),
  "0건으로 단정하거나 값을 지어내면 감점"),
]


def main():
    conn = psycopg2.connect(DSN)
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["예시유형", "예시질의", "llm 답변", "정답", "정답 SQL", "결측·판정기준"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF2F7")

    for typ, q, sql, fmt, note in QUESTIONS:
        if sql is None:
            answer, sql_text = fmt(None), "(SQL 없음 — 거부가 정답)"
        else:
            cur.execute(sql)
            rows = cur.fetchall()
            answer = fmt(rows) if rows else "(조회 0행 — '해당 데이터 없음'이 정답)"
            sql_text = " ".join(sql.split())          # 줄바꿈 정리해서 한 줄로
        ws.append([typ, q, "", answer, sql_text, note])
        print(f"{typ:12s} {q[:28]}...")

    for col, w in zip("ABCDEF", [13, 42, 30, 50, 55, 32]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    wb.save(OUT)
    conn.close()
    print(f"\n저장 완료: {OUT}")


if __name__ == "__main__":
    main()
